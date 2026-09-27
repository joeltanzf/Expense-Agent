from contextlib import closing
from pathlib import Path
import sqlite3, unittest, uuid
from unittest.mock import patch
import test_standalone as fixtures
from openpyxl import load_workbook
import agent


class YearAndDeletionTests(unittest.TestCase):
    setUp=fixtures.StandaloneTests.setUp
    tearDown=fixtures.StandaloneTests.tearDown
    add=fixtures.StandaloneTests.add
    pdf=fixtures.StandaloneTests.pdf
    upload=fixtures.StandaloneTests.upload
    count=fixtures.StandaloneTests.count

    def first_id(self):
        return self.db.execute('SELECT id FROM transactions ORDER BY rowid').fetchone()[0]

    def test_delete_manual_expense_zero_total_and_predelete_backup(self):
        self.add();entry=self.first_id()
        result=agent.delete_expense(self.db,entry)
        self.assertTrue(result['deleted']);self.assertTrue(result['synced']);self.assertEqual(self.count(),0)
        with closing(load_workbook(agent.workbook_path('2026'),data_only=True)) as wb:
            self.assertEqual(wb['September']['B4'].value,'Total:')
            self.assertEqual(wb['September']['C4'].value,0)
        snapshots=[]
        for p in (agent.DATA/'backups').glob('Database-*.sqlite3'):
            with closing(sqlite3.connect(p)) as backup:
                snapshots.append(backup.execute('SELECT COUNT(*) FROM transactions WHERE id=?',(entry,)).fetchone()[0])
        self.assertIn(1,snapshots)
        self.assertFalse(agent.delete_expense(self.db,entry)['deleted'])

    def test_deleted_pdf_reference_stays_excluded_on_overlaps(self):
        row=fixtures.ROW;path=self.pdf('first.pdf',[row]);self.upload(path)
        agent.delete_expense(self.db,self.first_id())
        self.assertEqual(self.upload(path)['added'],0)
        other=(*row[:3],'P999',*row[4:])
        result=self.upload(self.pdf('overlap.pdf',[row,other]))
        self.assertEqual((result['added'],result['duplicates'],self.count()),(1,1,1))
        self.assertEqual(self.db.execute('SELECT reference FROM transactions').fetchone()[0],'P999')

    def test_deleting_linked_manual_expense_also_excludes_pdf_reference(self):
        self.add(description='Custom description');manual=self.first_id()
        path=self.pdf('match.pdf',[fixtures.ROW]);review=self.upload(path);row=review['candidates'][0]
        self.upload(path,resolutions={row['id']:manual})
        agent.delete_expense(self.db,manual)
        result=self.upload(self.pdf('precision.pdf',[(*fixtures.ROW[:-1],'5.0000')]))
        self.assertEqual((result['added'],result['duplicates'],self.count()),(0,1,0))
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM deleted_transactions').fetchone()[0],2)

    def test_delete_while_year_file_is_locked_then_retry(self):
        self.add();entry=self.first_id();path=agent.workbook_path('2026');prior=path.read_bytes()
        with path.open('rb'):
            result=agent.delete_expense(self.db,entry)
            self.assertTrue(result['deleted']);self.assertFalse(result['synced']);self.assertEqual(self.count(),0)
            self.assertEqual(path.read_bytes(),prior)
            self.assertTrue(agent.status(self.db)['years'][0]['pending'])
        self.assertTrue(agent.export_workbook(self.db)['synced'])
        self.assertFalse(agent.status(self.db)['years'][0]['pending'])

    def test_separate_year_files_have_plain_month_names_and_separate_totals(self):
        self.add('2025-01-20','Past expense','2.00');self.add('2026-01-05','Current expense','9.50')
        for year,expected in [('2025',2),('2026',9.5)]:
            with closing(load_workbook(agent.workbook_path(year),data_only=True)) as wb:
                self.assertEqual(wb.sheetnames,['January'])
                self.assertEqual(wb['January']['A2'].value.year,int(year))
                self.assertEqual(wb['January']['C4'].value,expected)
        state=agent.status(self.db)
        self.assertEqual([y['year'] for y in state['years']],['2025','2026'])
        self.assertEqual([y['count'] for y in state['years']],[1,1])
        self.assertFalse((agent.DATA/'Expenses.xlsx').exists())

    def test_locked_old_year_does_not_stop_new_year(self):
        self.add('2025-01-01','Past','5.00');old=agent.workbook_path('2025')
        with old.open('rb'):
            self.add('2025-01-02','Pending in past year','3.00')
            self.add('2026-01-01','New year','7.00')
            years={y['year']:y for y in agent.status(self.db)['years']}
            self.assertTrue(years['2025']['pending']);self.assertFalse(years['2026']['pending'])
            with closing(load_workbook(agent.workbook_path('2026'),data_only=True)) as wb:
                self.assertEqual(wb['January']['C4'].value,7)
        self.assertTrue(agent.export_workbook(self.db)['synced'])

    def test_rebuild_selected_year_preserves_other_year_external_edits(self):
        for year in ['2025','2026']:
            self.add(year+'-01-01','Saved description','5.00')
            with closing(load_workbook(agent.workbook_path(year))) as wb:
                wb['January']['B2']='External '+year;wb.save(agent.workbook_path(year))
        other=agent.workbook_path('2026').read_bytes()
        result=agent.export_workbook(self.db,force=True,year='2025')
        self.assertFalse(result['synced']);self.assertEqual(agent.workbook_path('2026').read_bytes(),other)
        with closing(load_workbook(agent.workbook_path('2025'))) as wb:self.assertEqual(wb['January']['B2'].value,'Saved description')

    def test_upgrade_retains_old_combined_file_and_existing_database_records(self):
        entry='manual:'+str(uuid.uuid4())
        self.db.execute('INSERT INTO transactions VALUES (?,?,?,?,?,?,?,?)',(entry,'2025-01-01','Existing expense','5.00','Manual expense',entry,'Manual entry',0))
        agent.month_add(self.db,'2025-01-01')
        self.db.execute("INSERT INTO rules(pattern,replacement,mode) VALUES ('Existing','Old rule','contains')")
        self.db.execute('DROP TABLE year_exports');self.db.execute('DROP TABLE deleted_transactions');self.db.commit()
        legacy=agent.DATA/'Expenses.xlsx';legacy.write_bytes(b'fictional untouched older workbook')
        self.db.close();self.db=agent.connect()
        agent.export_workbook(self.db)
        self.assertEqual(legacy.read_bytes(),b'fictional untouched older workbook')
        self.assertEqual(self.count(),1)
        with closing(load_workbook(agent.workbook_path('2025'))) as wb:self.assertEqual(wb['January']['B2'].value,'Old rule')

    def test_stale_manual_retry_does_not_restore_deleted_expense(self):
        request=dict(date='2026-09-10',description='Test retry',price='5',entry_id=str(uuid.uuid4()))
        agent.add_expense(self.db,request);agent.delete_expense(self.db,self.first_id())
        with self.assertRaises(agent.AgentError):agent.add_expense(self.db,request)
        request['entry_id']=str(uuid.uuid4());agent.add_expense(self.db,request)
        self.assertEqual(self.count(),1)

    def test_backup_failure_prevents_deletion(self):
        self.add();entry=self.first_id()
        with patch.object(agent,'backup_database',side_effect=OSError('Simulated full disk')):
            with self.assertRaises(OSError):agent.delete_expense(self.db,entry)
        self.assertEqual(self.count(),1)

    def test_existing_unowned_year_file_is_backed_up_only_on_explicit_rebuild(self):
        path=agent.workbook_path('2026');path.write_bytes(b'other workbook')
        result=self.add()
        self.assertFalse(result['synced']);self.assertEqual(path.read_bytes(),b'other workbook')
        self.assertTrue(agent.export_workbook(self.db,force=True,year='2026')['synced'])
        self.assertTrue(any(p.read_bytes()==b'other workbook' for p in (agent.DATA/'backups').glob('*.xlsx')))

    def test_missing_selection_cannot_delete_anything(self):
        self.add()
        with self.assertRaises(agent.AgentError):agent.delete_expense(self.db,None)
        self.assertFalse(agent.delete_expense(self.db,'missing')['deleted']);self.assertEqual(self.count(),1)

    def test_workbook_backups_keep_newest_date_across_expense_years(self):
        self.add('2025-01-01', 'Older year', '5.00')
        folder = agent.DATA / 'backups'
        older = []
        for index in range(30):
            path = folder / f'Expenses-2026-20000101-000000-{index:06d}-fictional.xlsx'
            path.write_bytes(b'fictional workbook backup')
            older.append(path)
        previous = agent.workbook_path('2025').read_bytes()
        self.add('2025-01-02', 'New change to older year', '2.00')
        remaining = list(folder.glob('Expenses-*.xlsx'))
        self.assertEqual(len(remaining), 30)
        self.assertFalse(older[0].exists())
        self.assertTrue(older[-1].exists())
        self.assertTrue(any(p.name.startswith('Expenses-2025-') and
                            p.read_bytes() == previous for p in remaining))

if __name__=='__main__':unittest.main(verbosity=2)
