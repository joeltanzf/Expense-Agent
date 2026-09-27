"""Create ordinary Excel workbooks locally using openpyxl."""
from datetime import date
from decimal import Decimal
import textwrap
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties


def write_workbook(data, output_path):
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    currency = '"RM" #,##0.00;[Red]("RM" #,##0.00);"RM" 0.00'
    for month in data['months']:
        ws = wb.create_sheet(month['sheet'])
        rows = [r for r in data['transactions'] if r['date'].startswith(month['month'])]
        last = max(2, len(rows) + 1)
        total = last + 2
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = 'A2'
        for column, width in [('A', 18), ('B', 64), ('C', 19)]:
            ws.column_dimensions[column].width = width
        for cells in ws.iter_rows(min_row=1, max_row=total, max_col=3):
            ws.row_dimensions[cells[0].row].height = 23
            for cell in cells:
                cell.font = Font(name='Arial', size=11, color='263244')
                cell.alignment = Alignment(vertical='center')
        for col, title in enumerate(['Date', 'Description', 'Price'], 1):
            cell = ws.cell(1, col, title)
            cell.fill = PatternFill('solid', fgColor='18324F')
            cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        for index, row in enumerate(rows, 2):
            ws.cell(index, 1, date.fromisoformat(row['date']))
            description = ws.cell(index, 2, row['description'])
            # Merchant names are text even when they begin with '='.
            description.data_type = 's'
            description.alignment = Alignment(vertical='center', wrap_text=True)
            ws.cell(index, 3, float(Decimal(row['amount'])))
            if index % 2 == 0:
                for cell in ws[index][:3]:
                    cell.fill = PatternFill('solid', fgColor='F2F5F9')
            lines = sum(max(1, len(textwrap.wrap(part, width=55)))
                        for part in row['description'].split('\n'))
            ws.row_dimensions[index].height = max(23, lines * 16 + 7)
        for index in range(2, last + 1):
            ws.cell(index, 1).number_format = 'dd/mm/yyyy'
            ws.cell(index, 1).alignment = Alignment(horizontal='center', vertical='center')
        for index in range(2, total + 1):
            ws.cell(index, 3).number_format = currency
            ws.cell(index, 3).alignment = Alignment(horizontal='right', vertical='center')
        ws.cell(total, 2, 'Total:')
        ws.cell(total, 3, f'=SUM(C2:C{last})')
        for cell in ws[total][:3]:
            cell.fill = PatternFill('solid', fgColor='E1E9F3')
            cell.font = Font(name='Arial', size=11, bold=True, color='263244')
            cell.border = Border(top=Side(style='thin', color='18324F'))
        ws.row_dimensions[total].height = 29
    wb.save(output_path)
    wb.close()
    # Cache only the SUM formulas we just authored for non-calculating viewers.
    # Excel still recalculates when a price is edited.
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('', ns)
    with zipfile.ZipFile(output_path) as archive:
        parts = {name:archive.read(name) for name in archive.namelist()}
    # Match Excel's conventional font order, including older Open XML readers.
    styles = ET.fromstring(parts['xl/styles.xml'])
    order = ['b','i','strike','condense','extend','outline','shadow','u','vertAlign','sz','color','name','family','charset','scheme']
    for font in styles.findall('{' + ns + '}fonts/{' + ns + '}font'):
        font[:] = sorted(font, key=lambda node: order.index(node.tag.split('}')[-1]))
    parts['xl/styles.xml'] = ET.tostring(styles, encoding='utf-8', xml_declaration=True)
    for index, month in enumerate(data['months'], 1):
        name = f'xl/worksheets/sheet{index}.xml'
        root = ET.fromstring(parts[name])
        total = sum((Decimal(r['amount']) for r in data['transactions'] if r['date'].startswith(month['month'])), Decimal(0))
        formulas = root.findall('.//{' + ns + '}c[{' + ns + '}f]')
        if len(formulas) != 1:
            raise ValueError('Unexpected worksheet formulas.')
        cached = formulas[0].find('{' + ns + '}v')
        if cached is None:
            cached = ET.SubElement(formulas[0], '{' + ns + '}v')
        cached.text = str(total)
        parts[name] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    with zipfile.ZipFile(output_path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, raw in parts.items():
            archive.writestr(name, raw)
