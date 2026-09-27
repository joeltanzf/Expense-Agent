"""Regression checks with fictional PDFs and disposable data only."""
from contextlib import closing
from datetime import date
from decimal import Decimal
from pathlib import Path
import json, sqlite3, sys, tempfile, unittest, uuid
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'app'))
import agent
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pdfencrypt import StandardEncryption

PASSWORD='fictional-test-password'
ROW=('10/09/2026','Success','Payment','P001','Example Market','5.00')

class StandaloneTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.folder=Path(self.tmp.name)
        self.old=agent.DATA,agent.WORK
        agent.DATA,agent.WORK=self.folder/'data',self.folder/'work'
        self.db=agent.connect()
    def tearDown(self):
        self.db.close();agent.DATA,agent.WORK=self.old;self.tmp.cleanup()
    def pdf(self,name,rows):
        path=self.folder/name
        pdf=canvas.Canvas(str(path),pagesize=(842,595),invariant=True,
            encrypt=StandardEncryption(PASSWORD,ownerPassword='fictional-owner',strength=128))
        pdf.setFont('Helvetica',12);pdf.drawString(20,560,'TNG WALLET TRANSACTION HISTORY')
        pdf.setFont('Helvetica',8);pdf.drawString(20,535,'Wallet ID 000000000000')
        pdf.drawString(20,515,'FICTIONAL TEST DATA')
        for index,row in enumerate(rows):
            pdf.setFont('Helvetica',7)
            for x,value in zip([20,82,142,232,296,652],row):pdf.drawString(x,470-index*40,'RM '+value if x==652 else value)
        pdf.save();return path
    def upload(self,path,**kwargs):return agent.import_pdf(self.db,path,PASSWORD,**kwargs)
    def add(self,day='2026-09-10',description='Example Market',price='5.00'):
        return agent.add_expense(self.db,dict(date=day,description=description,price=price,entry_id=str(uuid.uuid4())))
    def count(self):return self.db.execute('SELECT COUNT(*) FROM transactions').fetchone()[0]
    def test_encrypted_pdf_overlap_expenses_only_and_one_time_password(self):
        rows=[ROW,('11/09/2026','Success','DuitNow QR','P002','Example Cafe','7.50'),
              ('11/09/2026','Success','Reload','R001','Reload','100.00'),
              ('11/09/2026','Success','Receive from Wallet','R002','Incoming','20.00'),
              ('11/09/2026','Success','GO+ Cash In','G001','Savings','10.00')]
        path=self.pdf('first.pdf',rows);first=self.upload(path);again=self.upload(path)
        overlap=self.upload(self.pdf('overlap.pdf',[ROW,('12/09/2026','Success','Payment','P003','Bus','2.00')]))
        self.assertEqual((first['added'],first['excluded'],again['added'],overlap['added'],overlap['duplicates'],self.count()),(2,3,0,1,1,3))
        self.assertFalse(agent.get_setting(self.db,'pdf_password'))
        for p in agent.DATA.rglob('*'):
            if p.is_file():self.assertNotIn(PASSWORD.encode(),p.read_bytes())
    def test_wrong_password_unknown_type_and_conflict_are_atomic(self):
        path=self.pdf('badpw.pdf',[ROW])
        with self.assertRaises(agent.AgentError):agent.import_pdf(self.db,path,'wrong')
        unknown=(*ROW[:2],'Unknown type',*ROW[3:])
        with self.assertRaises(agent.AgentError):self.upload(self.pdf('unknown.pdf',[ROW,unknown]))
        conflict=(*ROW[:-1],'15.00')
        with self.assertRaises(agent.AgentError):self.upload(self.pdf('conflict.pdf',[ROW,conflict]))
        self.assertEqual(self.count(),0)
    def test_duplicate_within_pdf_and_equivalent_decimals(self):
        first=self.upload(self.pdf('repeated.pdf',[ROW,ROW]))
        second=self.upload(self.pdf('precision.pdf',[(*ROW[:-1],'5.0000')]))
        self.assertEqual((first['added'],first['duplicates'],second['added'],second['duplicates'],self.count()),(1,1,0,1,1))
    def test_future_month_does_not_block_current_and_gap_fill(self):
        self.add();self.add('2026-12-10','Future','3.00')
        agent.ensure_months(self.db,date(2026,11,1));self.db.commit();agent.export_workbook(self.db)
        self.assertEqual([r[0] for r in self.db.execute('SELECT month FROM months ORDER BY month')],['2026-09','2026-10','2026-11','2026-12'])
    def test_manual_sort_formula_cache_and_year_names(self):
        self.add('2026-02-28','Later','0.20');self.add('2026-02-01','=Literal','0.10');self.add('2027-02-01','Next year','2')
        with closing(load_workbook(agent.workbook_path('2026'))) as wb:
            self.assertEqual(wb.sheetnames,['February'])
            self.assertEqual(wb['February']['A2'].value.date(),date(2026,2,1))
            self.assertEqual(wb['February']['B2'].data_type,'s')
            self.assertEqual(wb['February']['C5'].value,'=SUM(C2:C3)')
        with closing(load_workbook(agent.workbook_path('2026'),data_only=True)) as wb:self.assertEqual(wb['February']['C5'].value,0.3)
    def test_bad_rule_rejected_before_saving_and_database_backed_up(self):
        self.add()
        for bad in ['Bad\x01text','Bad\ud800','Bad\uffff']:
            with self.assertRaises(agent.AgentError):agent.dispatch(dict(action='add_rule',pattern='Example',replacement=bad))
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM rules').fetchone()[0],0)
        agent.dispatch(dict(action='add_rule',pattern='Example',replacement='Groceries'))
        backups=list((agent.DATA/'backups').glob('Database-*.sqlite3'))
        self.assertTrue(backups)
        with closing(sqlite3.connect(sorted(backups)[-1])) as backup:
            self.assertEqual(backup.execute('PRAGMA integrity_check').fetchone()[0],'ok')
            self.assertEqual(backup.execute('SELECT replacement FROM rules').fetchone()[0],'Groceries')
        with closing(load_workbook(agent.workbook_path('2026'))) as wb:self.assertEqual(wb['September']['B2'].value,'Groceries')
    def test_locked_workbook_retry_keeps_data_one_backup(self):
        self.add();original=(agent.workbook_path('2026')).read_bytes()
        with open(agent.workbook_path('2026'),'rb'):
            result=self.add('2026-09-11','Another','2')
            for _ in range(2):self.assertFalse(agent.export_workbook(self.db)['synced'])
            self.assertFalse(result['synced']);self.assertEqual(self.count(),2)
            identical=[p for p in (agent.DATA/'backups').glob('Expenses-*.xlsx') if p.read_bytes()==original]
            self.assertEqual(len(identical),1)
        self.assertTrue(agent.export_workbook(self.db)['synced'])
    def test_external_edit_is_preserved_until_forced_rebuild(self):
        self.add()
        with closing(load_workbook(agent.workbook_path('2026'))) as wb:
            wb['September']['B2']='External edit';wb.save(agent.workbook_path('2026'))
        self.assertEqual(agent.export_workbook(self.db)['reason'],'workbook_modified')
        self.assertTrue(agent.export_workbook(self.db,force=True)['synced'])
    def test_manual_pdf_match_review_link_and_overlap(self):
        self.add(description='My groceries')
        path=self.pdf('match.pdf',[ROW]);review=self.upload(path)
        self.assertTrue(review['needs_review']);self.assertEqual(self.count(),1)
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM imports').fetchone()[0],0)
        row=review['candidates'][0];decisions={row['id']:row['matches'][0]['id']}
        result=self.upload(path,resolutions=decisions,expected_hash=review['pdf_hash'])
        self.assertEqual((result['added'],result['duplicates'],self.count()),(0,1,1))
        repeated=self.upload(self.pdf('overlap.pdf',[(*ROW[:-1],'5.0000')]))
        self.assertEqual(repeated['added'],0)
        self.assertEqual(self.db.execute('SELECT description FROM transactions').fetchone()[0],'My groceries')
    def test_review_keep_both_and_changed_pdf_guard(self):
        self.add();path=self.pdf('match.pdf',[ROW]);review=self.upload(path)
        with self.assertRaises(agent.AgentError):self.upload(path,expected_hash='changed')
        result=self.upload(path,resolutions={review['candidates'][0]['id']:'keep'},expected_hash=review['pdf_hash'])
        self.assertEqual((result['added'],self.count()),(1,2))
    def test_one_manual_cannot_link_to_two_pdf_rows(self):
        self.add();path=self.pdf('twomatches.pdf',[ROW,(*ROW[:3],'P002',*ROW[4:])]);review=self.upload(path)
        decisions={r['id']:r['matches'][0]['id'] for r in review['candidates']}
        with self.assertRaises(agent.AgentError):self.upload(path,resolutions=decisions)
        self.assertEqual(self.count(),1)
    def test_dpapi_failure_never_stores_plaintext_and_once_still_works(self):
        with patch.object(agent,'protect',side_effect=agent.AgentError('Windows encryption unavailable')):
            with self.assertRaises(agent.AgentError):agent.dispatch(dict(action='save_settings',pdf_password=PASSWORD))
            self.assertFalse(agent.get_setting(self.db,'pdf_password'))
            self.assertEqual(self.upload(self.pdf('once.pdf',[ROW]))['added'],1)
    def test_export_failure_preserves_previous_workbook_and_saved_expense(self):
        self.add();previous=(agent.workbook_path('2026')).read_bytes()
        with patch.object(agent,'write_workbook',side_effect=ValueError('simulated failure')):
            result=self.add('2026-09-12','Still saved','2')
        self.assertTrue(result['saved']);self.assertFalse(result['synced']);self.assertEqual(self.count(),2)
        self.assertEqual((agent.workbook_path('2026')).read_bytes(),previous)
        self.assertTrue(agent.export_workbook(self.db)['synced'])

if __name__=='__main__':unittest.main(verbosity=2)
